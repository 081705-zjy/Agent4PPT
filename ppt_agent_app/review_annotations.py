"""Strip review annotation textboxes from a draft pptx.

The annotations are tagged with `shape.name == "REVIEW_ANNOTATION_v1"` at
inject time (DRAFT_ASSEMBLY phase J). This module removes them so the deck
can be shared/downstream-rendered without the red-border review overlay.

Usage:
    python3 -m ppt_agent_app.review_annotations strip <input.pptx> [output.pptx]

If output.pptx is omitted, the input file is overwritten in place.
"""

from __future__ import annotations

import sys
from pathlib import Path


REVIEW_ANNOTATION_NAME = "REVIEW_ANNOTATION_v1"


def strip(input_path: str | Path, output_path: str | Path | None = None) -> dict:
    """Remove all shapes whose name == REVIEW_ANNOTATION_v1 from each slide.

    Returns ``{"ok": True, "removed": N, "slides_touched": M, "output": path}``.
    """
    from pptx import Presentation

    in_path = Path(input_path)
    if not in_path.exists():
        return {"ok": False, "error": f"input not found: {in_path}"}

    out_path = Path(output_path) if output_path else in_path
    pres = Presentation(str(in_path))

    removed = 0
    slides_touched = 0
    for slide in pres.slides:
        shapes_to_remove = [sh for sh in slide.shapes
                            if getattr(sh, "name", "") == REVIEW_ANNOTATION_NAME]
        if not shapes_to_remove:
            continue
        slides_touched += 1
        # python-pptx doesn't expose a direct shape.remove(); drop via spTree.
        sp_tree = slide.shapes._spTree
        for sh in shapes_to_remove:
            sp_tree.remove(sh._element)
            removed += 1

    pres.save(str(out_path))
    return {"ok": True, "removed": removed, "slides_touched": slides_touched,
            "input": str(in_path), "output": str(out_path)}


def _main(argv: list[str]) -> int:
    if len(argv) < 2 or argv[0] != "strip":
        print(__doc__.strip())
        return 1
    in_path = argv[1]
    out_path = argv[2] if len(argv) > 2 else None
    result = strip(in_path, out_path)
    if not result.get("ok"):
        print(f"FAIL: {result.get('error')}")
        return 2
    print(f"removed {result['removed']} annotations across "
          f"{result['slides_touched']} slides → {result['output']}")
    return 0


if __name__ == "__main__":
    sys.exit(_main(sys.argv[1:]))
